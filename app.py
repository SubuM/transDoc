import streamlit as st
from pdf2image import convert_from_bytes
import pytesseract
from langdetect import detect
from langdetect.lang_detect_exception import LangDetectException
from PIL import Image
import io
import fitz
import os
import tempfile
from docx import Document

try:
    import argostranslate.package
    import argostranslate.translate
    HAS_ARGOS = True
except Exception:
    HAS_ARGOS = False

st.set_page_config(page_title="transDoc", layout="wide")

st.title("Document Translator")
st.write("Upload PDF, DOC, or DOCX. OCR, translate to English offline. Download as PDF or DOCX.")

@st.cache_data
def translate_text_offline(text, from_code):
    if from_code == 'en' or not text.strip():
        return text
    try:
        if HAS_ARGOS:
            translated = argostranslate.translate.translate(text, from_code, 'en')
            if translated:
                return translated
    except Exception:
        pass
    return text

if 'processed' not in st.session_state:
    st.session_state.processed = False
    st.session_state.translated_texts = []
    st.session_state.images = []
    st.session_state.detected_lang = 'en'
    st.session_state.filename = ''

uploaded_file = st.file_uploader("Upload file", type=["pdf", "doc", "docx"], key="file_uploader")

if uploaded_file is not None and not st.session_state.processed:
    if st.button("Translate File", type="primary"):
        with st.spinner("Processing..."):
            try:
                file_bytes = uploaded_file.read()
                filename = uploaded_file.name.lower()
                st.session_state.filename = uploaded_file.name
                
                all_images = []
                detected_lang = "en"
                translated_content = []
                
                if filename.endswith('.pdf'):
                    pages = convert_from_bytes(file_bytes, dpi=200)
                    all_images = pages
                    try:
                        sample_text = pytesseract.image_to_string(pages[0])
                        if sample_text.strip():
                            try:
                                detected_lang = detect(sample_text)
                            except LangDetectException:
                                pass
                    except Exception:
                        pass
                    
                    for page_num, img in enumerate(pages):
                        try:
                            ocr_data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
                            n_boxes = len(ocr_data['text'])
                            page_texts = []
                            for i in range(n_boxes):
                                text = ocr_data['text'][i].strip()
                                if int(ocr_data['conf'][i]) > 40 and text and len(text) > 1:
                                    translated = translate_text_offline(text, detected_lang)
                                    if translated and translated.strip():
                                        page_texts.append(translated)
                            translated_content.append(' '.join(page_texts))
                        except Exception:
                            translated_content.append('')
                
                elif filename.endswith('.docx'):
                    try:
                        doc = Document(io.BytesIO(file_bytes))
                    except Exception as e:
                        st.error(f"Error reading DOCX: {e}")
                        st.stop()
                    full_text = []
                    for para in doc.paragraphs:
                        full_text.append(para.text)
                    combined = '\n'.join(full_text)
                    try:
                        if combined.strip():
                            detected_lang = detect(combined[:1000])
                    except Exception:
                        pass
                    for para in doc.paragraphs:
                        if para.text.strip():
                            translated = translate_text_offline(para.text, detected_lang)
                            translated_content.append(translated)
                        else:
                            translated_content.append('')
                    # Extract images safely
                    try:
                        for rel in doc.part.rels.values():
                            try:
                                if "image" in rel.target_ref:
                                    img_data = rel.target_part.blob
                                    img = Image.open(io.BytesIO(img_data))
                                    # Convert to RGB if needed
                                    if img.mode in ('RGBA', 'P'):
                                        img = img.convert('RGB')
                                    all_images.append(img)
                            except Exception:
                                continue
                    except Exception:
                        pass
                
                elif filename.endswith('.doc'):
                    try:
                        with tempfile.NamedTemporaryFile(delete=False, suffix='.doc') as tmp:
                            tmp.write(file_bytes)
                            tmp_path = tmp.name
                        pdf_path = tmp_path.replace('.doc', '.pdf')
                        os.system(f'soffice --headless --convert-to pdf "{tmp_path}" --outdir "{os.path.dirname(tmp_path)}" >/dev/null 2>&1')
                        if os.path.exists(pdf_path):
                            with open(pdf_path, 'rb') as f:
                                pdf_bytes_conv = f.read()
                            pages = convert_from_bytes(pdf_bytes_conv, dpi=200)
                            all_images = pages
                            try:
                                sample_text = pytesseract.image_to_string(pages[0])
                                if sample_text.strip():
                                    try:
                                        detected_lang = detect(sample_text)
                                    except LangDetectException:
                                        pass
                            except Exception:
                                pass
                            for page_num, img in enumerate(pages):
                                try:
                                    ocr_data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)
                                    n_boxes = len(ocr_data['text'])
                                    page_texts = []
                                    for i in range(n_boxes):
                                        text = ocr_data['text'][i].strip()
                                        if int(ocr_data['conf'][i]) > 40 and text and len(text) > 1:
                                            translated = translate_text_offline(text, detected_lang)
                                            if translated and translated.strip():
                                                page_texts.append(translated)
                                    translated_content.append(' '.join(page_texts))
                                except Exception:
                                    translated_content.append('')
                            os.unlink(pdf_path)
                        os.unlink(tmp_path)
                    except Exception as e:
                        st.error(f"Error: {e}")
                
                st.session_state.translated_texts = translated_content
                st.session_state.images = all_images
                st.session_state.detected_lang = detected_lang
                st.session_state.processed = True
                st.rerun()
            except Exception as e:
                st.error(f"Error: {e}")
                st.exception(e)

if st.session_state.processed:
    st.success(f"Translated from {st.session_state.detected_lang} to English (offline)")
    
    format_choice = st.radio("Download format", ["PDF", "DOCX"], horizontal=True)
    
    if format_choice == "PDF":
        new_doc = fitz.open()
        if st.session_state.images:
            for idx, img in enumerate(st.session_state.images):
                try:
                    if img.mode != 'RGB':
                        img = img.convert('RGB')
                    img_byte_arr = io.BytesIO()
                    img.save(img_byte_arr, format='PNG')
                    img_bytes = img_byte_arr.getvalue()
                    page_width = img.width * 0.75
                    page_height = img.height * 0.75
                    new_page = new_doc.new_page(width=page_width, height=page_height)
                    new_page.insert_image(fitz.Rect(0, 0, page_width, page_height), stream=img_bytes)
                    if idx < len(st.session_state.translated_texts):
                        text = st.session_state.translated_texts[idx]
                        if text.strip():
                            try:
                                new_page.insert_textbox(fitz.Rect(50, page_height * 0.1, page_width - 50, page_height * 0.9), 
                                                       text, fontsize=9, fontname="helv")
                            except Exception:
                                pass
                except Exception:
                    continue
        else:
            new_page = new_doc.new_page()
            all_text = '\n\n'.join(st.session_state.translated_texts)
            try:
                new_page.insert_textbox(fitz.Rect(50, 50, 550, 750), all_text, fontsize=11, fontname="helv")
            except Exception:
                pass
        
        pdf_output = io.BytesIO()
        new_doc.save(pdf_output)
        pdf_output.seek(0)
        new_doc.close()
        
        st.download_button(
            label="Download PDF",
            data=pdf_output.getvalue(),
            file_name=f"translated_{st.session_state.filename.rsplit('.',1)[0]}.pdf",
            mime="application/pdf"
        )
    
    elif format_choice == "DOCX":
        doc = Document()
        if st.session_state.translated_texts:
            for idx, text in enumerate(st.session_state.translated_texts):
                if text.strip():
                    doc.add_paragraph(text)
                if idx < len(st.session_state.translated_texts) - 1 and text.strip():
                    doc.add_page_break()
        doc_output = io.BytesIO()
        doc.save(doc_output)
        doc_output.seek(0)
        
        st.download_button(
            label="Download DOCX",
            data=doc_output.getvalue(),
            file_name=f"translated_{st.session_state.filename.rsplit('.',1)[0]}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    
    if st.button("Translate Another File"):
        st.session_state.processed = False
        st.session_state.translated_texts = []
        st.session_state.images = []
        st.rerun()
